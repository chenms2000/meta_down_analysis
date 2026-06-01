"""Run the European expanded/soft seed evidence and upgrade pipeline.

The pipeline is conservative and source-bound:

1. Export the current expanded candidate queue from the service precheck.
2. Collect evidence from the European Nature Genetics supplementary tables,
   current graph/index support, and existing review/annotation files.
3. Generate a decision table for strict/soft/analog/class/ratio/unresolved.
4. Update the European identity review with only high-confidence strict
   decisions and conservative class/ratio triage rows.
5. Extend the compound identity overlay for strict PubChem CIDs that are still
   absent from the local compound index.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import re
import sys
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

try:
    import openpyxl
except Exception:  # pragma: no cover - handled in main.
    openpyxl = None

try:
    import pyarrow.dataset as ds
except Exception:  # pragma: no cover
    ds = None


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from build_graph_projection import normalize_id_token, normalize_lookup_key  # noqa: E402


DEFAULT_RELEASE_ID = "mvp_20260513T002254"
DEFAULT_SUPPLEMENT_URL = (
    "https://static-content.springer.com/esm/art%3A10.1038%2Fs41588-022-01270-1/"
    "MediaObjects/41588_2022_1270_MOESM4_ESM.xlsx"
)
DEFAULT_SUPPLEMENT_PATH = "raw_lake/European/source_supplements/41588_2022_1270_MOESM4_ESM.xlsx"
DEFAULT_OUTPUT_DIR = "manual_sources/european_expanded_identity"
DEFAULT_REVIEW = "manual_sources/european_trait_identity/european_trait_identity_review.csv"
DEFAULT_OVERLAY_ROOT = "manual_sources/compound_identity_overlays"
ACCEPTED_REVIEW_STATUSES = {"reviewed", "auto_reviewed", "accepted", "confirmed"}
STRICT_EVIDENCE_SOURCE = "PMID:36635386 Supplementary Table 2; auto rule european_expanded_supplement_v1"


def require_openpyxl() -> None:
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to read the European supplementary XLSX.")


def read_csv_rows(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    if not path.exists():
        return [], []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return [dict(row) for row in reader], list(reader.fieldnames or [])


def write_csv_rows(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def split_values(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text or text.upper() in {"NA", "N/A", "NONE", "NULL"}:
        return []
    parts = re.split(r"[|;]", text)
    out: list[str] = []
    seen: set[str] = set()
    for part in parts:
        cleaned = str(part or "").strip()
        if not cleaned or cleaned.upper() in {"NA", "N/A", "NONE", "NULL"}:
            continue
        key = cleaned.casefold()
        if key not in seen:
            seen.add(key)
            out.append(cleaned)
    return out


def join_values(values: list[Any]) -> str:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        for part in split_values(value):
            key = part.casefold()
            if key not in seen:
                seen.add(key)
                out.append(part)
    return "|".join(out)


def normalize_trait_name(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\b(?:levels?|measurements?)\b", " ", text, flags=re.I)
    text = text.replace("α", "alpha").replace("β", "beta").replace("ω", "omega")
    text = re.sub(r"[^A-Za-z0-9]+", " ", text.casefold())
    return re.sub(r"\s+", " ", text).strip()


def risk_reason(value: Any) -> str:
    text = str(value or "").strip()
    low = text.casefold()
    if re.search(r"\bX[- ]?\d+\b", text, flags=re.I):
        return "platform_x_code"
    if re.search(r"\bratio\b|\s+to\s+", low):
        return "ratio_text"
    if "*" in text:
        return "platform_inferred_identity"
    if " / " in text or re.search(r"\s/[^\s]|[^\s]/\s", text):
        return "slash_ratio_or_lipid"
    if re.search(r"\bor\b|\[[^\]]+\]", low):
        return "isomer_or_positional_pool"
    if re.search(
        r"\b(?:pc|pe|pi|ps|tg|dg|cer|sm)\b(?:\s|\(|$)|"
        r"sphingomyelin.*[,/]|ceramide.*[,/]|acylcarnitine\s+pool",
        low,
    ):
        return "lipid_or_pool_text"
    return ""


def load_script_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def download_supplement(path: Path, url: str) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(url, headers={"User-Agent": "metabo-expanded-identity-pipeline/1.0"})
    with urllib.request.urlopen(request, timeout=60) as response:
        path.write_bytes(response.read())


def row_dict_from_sheet(sheet, header_row: int) -> list[dict[str, str]]:
    headers = [str(value or "").strip() for value in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        row = {headers[index]: str(value or "").strip() for index, value in enumerate(values) if index < len(headers)}
        if any(row.values()):
            rows.append(row)
    return rows


def load_supplement(path: Path) -> dict[str, Any]:
    require_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    st2_rows = row_dict_from_sheet(wb["ST2"], 3)
    st13_rows = row_dict_from_sheet(wb["ST13"], 3)
    metabolite_by_name: dict[str, dict[str, str]] = {}
    for row in st2_rows:
        name = row.get("Metabolites", "")
        key = normalize_trait_name(name)
        if key:
            metabolite_by_name.setdefault(key, row)
    ratios_by_name: dict[str, dict[str, str]] = {}
    for row in st13_rows:
        ratio = row.get("Metabolite ratios", "")
        key = normalize_trait_name(ratio)
        if key:
            ratios_by_name.setdefault(key, row)
    return {
        "metabolite_by_name": metabolite_by_name,
        "ratio_by_name": ratios_by_name,
        "st2_count": len(st2_rows),
        "st13_count": len(st13_rows),
    }


def first_value(row: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = str(row.get(key) or "").strip()
        if value and value.upper() not in {"NA", "N/A"}:
            return value
    return ""


def supplement_identifiers(row: dict[str, str]) -> dict[str, str]:
    return {
        "pubchem_cid": first_value(row, "PUBCHEM_curated", "PUBCHEM"),
        "hmdb_id": first_value(row, "HMDB_curated", "HMDB"),
        "kegg_id": first_value(row, "KEGG"),
        "inchikey": first_value(row, "INCHIKEY"),
        "formula": "",
        "super_pathway": first_value(row, "SUPER_PATHWAY"),
        "sub_pathway": first_value(row, "SUB_PATHWAY"),
        "metabolon_compound_id": first_value(row, "Compound ID (Metabolon)"),
        "metabolon_type": first_value(row, "TYPE"),
        "platform": first_value(row, "PLATFORM"),
    }


def local_pubchem_cids(compound_dir: Path) -> set[str]:
    if ds is None:
        return set()
    path = compound_dir / "compound_identifier_index.parquet"
    if not path.exists():
        return set()
    table = ds.dataset(str(path), format="parquet").to_table(columns=["namespace", "lookup_key"])
    out: set[str] = set()
    for namespace, lookup_key in zip(table.column("namespace").to_pylist(), table.column("lookup_key").to_pylist()):
        if str(namespace or "").upper() == "CID":
            out.add(normalize_id_token(lookup_key))
    return out


def load_european_records(path: Path) -> list[dict[str, str]]:
    rows, _fields = read_csv_rows(path)
    return rows


def build_service(workspace: Path, release_id: str):
    metabo_service = load_script_module("metabo_service_for_expanded_pipeline", workspace / "scripts" / "metabo_service.py")
    return metabo_service.MetaboService(
        workspace=workspace,
        normalized_root=workspace / "normalized_store",
        graph_root=workspace / "graph_projection",
        pubchem_root=workspace / "pubchem_cid_cache",
        compound_root=workspace / "compound_match_index",
        literature_root=workspace / "literature_resolved",
        prediction_root=workspace / "manual_sources" / "prediction_overlays",
        release_id=release_id,
    )


def source_trait_text(row: dict[str, Any]) -> str:
    record = row.get("record") or {}
    return first_value(record, "reported_trait", "reportedTrait", "trait", "name")


def accession_for_row(row: dict[str, Any]) -> str:
    record = row.get("record") or {}
    return first_value(record, "accession_id", "trait", "id", "input_id") or row.get("input_id", "")


def top_candidate(row: dict[str, Any]) -> dict[str, Any]:
    candidates = row.get("resolution", {}).get("candidates") or []
    return candidates[0] if candidates else {}


def candidate_xrefs(candidate: dict[str, Any]) -> str:
    return join_values(candidate.get("external_xrefs") or [])


def candidate_match_sources(candidate: dict[str, Any]) -> str:
    return join_values(
        [
            f"{match.get('source_table', '')}:{match.get('match_field', '')}:{match.get('raw_value', '')}"
            for match in candidate.get("matches", [])[:6]
        ]
    )


def collect_queue(precheck: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    grouped: dict[str, dict[str, Any]] = {}
    candidate_rows: list[dict[str, Any]] = []
    source_rows = {
        row.get("input_id", ""): row
        for row in [*precheck.get("ambiguous", []), *precheck.get("unmatched", []), *precheck.get("invalid", [])]
    }
    for expanded in precheck.get("expanded_candidates", []) or []:
        meta = expanded.get("expanded_candidate") or {}
        original_input_id = meta.get("original_input_id") or expanded.get("input_id", "").split("::", 1)[0]
        source = source_rows.get(original_input_id, expanded)
        accession = accession_for_row(source)
        trait_text = source_trait_text(source)
        candidate = top_candidate(expanded)
        group = grouped.setdefault(
            original_input_id,
            {
                "original_input_id": original_input_id,
                "accession_id": accession,
                "reported_trait": trait_text,
                "current_seed_class": meta.get("seed_class", ""),
                "current_reason": meta.get("reason", ""),
                "source_status": meta.get("source_status", ""),
                "top_score": source.get("resolution", {}).get("top_score", 0.0),
                "top_margin": source.get("resolution", {}).get("top_margin", 0.0),
                "candidate_count": 0,
                "candidate_names": [],
                "candidate_uids": [],
                "candidate_xrefs": [],
                "record": source.get("record", {}),
            },
        )
        group["candidate_count"] += 1
        group["candidate_names"].append(candidate.get("display_name") or candidate.get("entity_uid") or "")
        group["candidate_uids"].append(candidate.get("entity_uid") or "")
        group["candidate_xrefs"].append(candidate_xrefs(candidate))
        candidate_rows.append(
            {
                "original_input_id": original_input_id,
                "accession_id": accession,
                "reported_trait": trait_text,
                "seed_class": meta.get("seed_class", ""),
                "candidate_uid": candidate.get("entity_uid", ""),
                "candidate_name": candidate.get("display_name", ""),
                "candidate_score": candidate.get("score", ""),
                "candidate_xrefs": candidate_xrefs(candidate),
                "candidate_match_sources": candidate_match_sources(candidate),
                "weight_multiplier": meta.get("weight_multiplier", ""),
                "reason": meta.get("reason", ""),
            }
        )
    for group in grouped.values():
        group["candidate_names"] = join_values(group["candidate_names"])
        group["candidate_uids"] = join_values(group["candidate_uids"])
        group["candidate_xrefs"] = join_values(group["candidate_xrefs"])
    return candidate_rows, grouped


def decide_for_group(group: dict[str, Any], supplement: dict[str, Any], supplement_url: str, local_cids: set[str]) -> dict[str, Any]:
    trait = str(group.get("reported_trait") or "")
    record = group.get("record") or {}
    candidate_names = split_values(group.get("candidate_names"))
    search_names = [
        trait,
        first_value(record, "name", "metabolite", "compound"),
        *split_values(first_value(record, "mapped_names", "candidate_names")),
        *candidate_names,
    ]
    search_names = [name for name in search_names if name]
    exact_supplement = None
    exact_name = ""
    for name in search_names:
        key = normalize_trait_name(name)
        if key in supplement["metabolite_by_name"]:
            exact_supplement = supplement["metabolite_by_name"][key]
            exact_name = name
            break
    ratio_supplement = None
    for name in search_names:
        key = normalize_trait_name(name)
        if key in supplement["ratio_by_name"]:
            ratio_supplement = supplement["ratio_by_name"][key]
            break
    exact_supplement_name = first_value(exact_supplement or {}, "Metabolites") if exact_supplement else ""
    combined_text = " | ".join([part for part in [trait, exact_supplement_name] if part])
    if not combined_text:
        combined_text = " | ".join([trait, group.get("candidate_names", "")])
    risk = risk_reason(combined_text)
    identifiers = supplement_identifiers(exact_supplement or {}) if exact_supplement else {}
    has_stable_id = bool(
        identifiers.get("pubchem_cid")
        or identifiers.get("hmdb_id")
        or identifiers.get("kegg_id")
        or identifiers.get("inchikey")
    )
    decision = {
        "original_input_id": group.get("original_input_id", ""),
        "accession_id": group.get("accession_id", ""),
        "reported_trait": trait,
        "current_seed_class": group.get("current_seed_class", ""),
        "candidate_count": group.get("candidate_count", 0),
        "candidate_names": group.get("candidate_names", ""),
        "candidate_uids": group.get("candidate_uids", ""),
        "top_score": group.get("top_score", ""),
        "top_margin": group.get("top_margin", ""),
        "decision_class": group.get("current_seed_class", ""),
        "decision_confidence": "low",
        "upgrade_action": "none",
        "decision_reason": "",
        "review_status": "",
        "reviewed_name": "",
        "reviewed_pubchem_cid": "",
        "reviewed_hmdb_id": "",
        "reviewed_chebi_id": "",
        "reviewed_kegg_id": "",
        "reviewed_inchikey": "",
        "supplement_metabolon_id": "",
        "supplement_super_pathway": "",
        "supplement_sub_pathway": "",
        "supplement_type": "",
        "local_cid_status": "",
        "evidence_source": "",
        "evidence_url": "",
        "search_urls": suggested_search_urls(trait, identifiers),
    }
    if ratio_supplement or group.get("current_seed_class") == "ratio_component" or risk in {"ratio_text"}:
        decision.update(
            {
                "decision_class": "ratio_component",
                "decision_confidence": "high" if ratio_supplement else "medium",
                "upgrade_action": "triage_ratio_component",
                "decision_reason": "ratio_or_composite_trait",
                "review_status": "auto_triaged",
                "evidence_source": "PMID:36635386 Supplementary Table 13" if ratio_supplement else "deterministic ratio text triage",
                "evidence_url": supplement_url,
            }
        )
        if ratio_supplement:
            decision["reviewed_name"] = first_value(ratio_supplement, "Metabolite ratios")
            decision["supplement_super_pathway"] = join_values([ratio_supplement.get("metabolite 1"), ratio_supplement.get("metabolite 2")])
        return decision
    if risk:
        decision.update(
            {
                "decision_class": "class_or_pool",
                "decision_confidence": "high" if exact_supplement else "medium",
                "upgrade_action": "triage_class_or_pool",
                "decision_reason": risk,
                "review_status": "auto_triaged",
                "reviewed_name": first_value(exact_supplement or {}, "Metabolites") or trait,
                "evidence_source": "PMID:36635386 Supplementary Table 2" if exact_supplement else "deterministic lipid/pool text triage",
                "evidence_url": supplement_url,
            }
        )
        if exact_supplement:
            decision.update(
                {
                    "supplement_metabolon_id": identifiers.get("metabolon_compound_id", ""),
                    "supplement_super_pathway": identifiers.get("super_pathway", ""),
                    "supplement_sub_pathway": identifiers.get("sub_pathway", ""),
                    "supplement_type": identifiers.get("metabolon_type", ""),
                    "reviewed_pubchem_cid": identifiers.get("pubchem_cid", ""),
                    "reviewed_hmdb_id": identifiers.get("hmdb_id", ""),
                    "reviewed_kegg_id": identifiers.get("kegg_id", ""),
                    "reviewed_inchikey": identifiers.get("inchikey", ""),
                }
            )
        return decision
    if exact_supplement and has_stable_id and identifiers.get("metabolon_type", "").upper() == "NAMED":
        cid = normalize_id_token(identifiers.get("pubchem_cid", ""))
        decision.update(
            {
                "decision_class": "strict_identity",
                "decision_confidence": "high",
                "upgrade_action": "promote_strict_identity",
                "decision_reason": "exact_non_pool_name_in_source_supplement_with_stable_identifier",
                "review_status": "auto_reviewed",
                "reviewed_name": first_value(exact_supplement, "Metabolites") or exact_name,
                "reviewed_pubchem_cid": identifiers.get("pubchem_cid", ""),
                "reviewed_hmdb_id": identifiers.get("hmdb_id", ""),
                "reviewed_kegg_id": identifiers.get("kegg_id", ""),
                "reviewed_inchikey": identifiers.get("inchikey", ""),
                "supplement_metabolon_id": identifiers.get("metabolon_compound_id", ""),
                "supplement_super_pathway": identifiers.get("super_pathway", ""),
                "supplement_sub_pathway": identifiers.get("sub_pathway", ""),
                "supplement_type": identifiers.get("metabolon_type", ""),
                "local_cid_status": "local_cid_present" if cid and cid in local_cids else ("local_cid_absent" if cid else ""),
                "evidence_source": STRICT_EVIDENCE_SOURCE,
                "evidence_url": supplement_url,
            }
        )
        return decision
    if exact_supplement:
        decision.update(
            {
                "decision_class": "soft_identity",
                "decision_confidence": "medium",
                "upgrade_action": "retain_soft_identity",
                "decision_reason": "source_supplement_name_match_without_full_strict_conditions",
                "reviewed_name": first_value(exact_supplement, "Metabolites"),
                "reviewed_pubchem_cid": identifiers.get("pubchem_cid", ""),
                "reviewed_hmdb_id": identifiers.get("hmdb_id", ""),
                "reviewed_kegg_id": identifiers.get("kegg_id", ""),
                "reviewed_inchikey": identifiers.get("inchikey", ""),
                "supplement_metabolon_id": identifiers.get("metabolon_compound_id", ""),
                "supplement_super_pathway": identifiers.get("super_pathway", ""),
                "supplement_sub_pathway": identifiers.get("sub_pathway", ""),
                "supplement_type": identifiers.get("metabolon_type", ""),
                "evidence_source": "PMID:36635386 Supplementary Table 2",
                "evidence_url": supplement_url,
            }
        )
        return decision
    current = group.get("current_seed_class", "")
    decision.update(
        {
            "decision_class": current or "unresolved",
            "decision_confidence": "low" if current != "unresolved" else "none",
            "upgrade_action": "retain_current_class",
            "decision_reason": "no_source_supplement_exact_match",
        }
    )
    return decision


def suggested_search_urls(trait: str, identifiers: dict[str, str]) -> str:
    encoded = re.sub(r"\s+", "%20", str(trait or "").strip())
    urls = [
        f"https://hmdb.ca/textquery?utf8=%E2%9C%93&query={encoded}",
        f"https://rest.kegg.jp/find/compound/{encoded}",
        f"https://www.metabolomicsworkbench.org/databases/refmet/refmet_search.php?refmet_name={encoded}",
    ]
    cid = identifiers.get("pubchem_cid", "")
    inchikey = identifiers.get("inchikey", "")
    if cid:
        urls.append(f"https://pubchem.ncbi.nlm.nih.gov/compound/{cid}")
        urls.append(f"https://www.lipidmaps.org/rest/compound/pubchem_cid/{cid}/classification/json")
    if inchikey:
        urls.append(f"http://classyfire.wishartlab.com/entities/{inchikey}.json")
        urls.append(f"http://mychem.info/v1/chem/{inchikey}")
    return "|".join(urls)


def review_row_key(row: dict[str, Any]) -> tuple[str, str, str]:
    accession = str(row.get("accession_id") or "").strip()
    name = normalize_trait_name(row.get("reviewed_name") or row.get("component_name") or row.get("reported_trait"))
    scope = str(row.get("identity_scope") or "").strip()
    return accession, name, scope


def update_identity_review(review_path: Path, decisions: list[dict[str, Any]]) -> dict[str, int]:
    review_rows, fieldnames = read_csv_rows(review_path)
    required_fields = [
        "priority",
        "accession_id",
        "reported_trait",
        "source_resolution_status",
        "review_bucket",
        "suggested_identity_scope",
        "suggested_action",
        "component_index",
        "component_name",
        "relationship_to_trait",
        "candidate_names",
        "mapped_names",
        "pubchem_cids",
        "pubchem_titles",
        "review_status",
        "identity_scope",
        "reviewed_name",
        "reviewed_pubchem_cid",
        "reviewed_hmdb_id",
        "reviewed_chebi_id",
        "reviewed_inchikey",
        "evidence_source",
        "evidence_url",
        "notes",
        "summary_statistics_url",
        "pubmed_id",
        "paper_title",
        "autofill_decision",
        "autofill_confidence",
        "local_cid_status",
    ]
    for field in required_fields:
        if field not in fieldnames:
            fieldnames.append(field)
    by_key = {review_row_key(row): row for row in review_rows}
    added = updated = skipped = 0
    for decision in decisions:
        action = decision.get("upgrade_action", "")
        if action not in {"promote_strict_identity", "triage_ratio_component", "triage_class_or_pool"}:
            skipped += 1
            continue
        scope = decision["decision_class"]
        reviewed_name = decision.get("reviewed_name") or decision.get("reported_trait") or ""
        key = (decision.get("accession_id", ""), normalize_trait_name(reviewed_name), scope)
        row = by_key.get(key)
        if row is None:
            row = {field: "" for field in fieldnames}
            review_rows.append(row)
            by_key[key] = row
            added += 1
        else:
            updated += 1
        row.update(
            {
                "priority": row.get("priority") or ("10" if scope == "strict_identity" else "30"),
                "accession_id": decision.get("accession_id", ""),
                "reported_trait": decision.get("reported_trait", ""),
                "source_resolution_status": decision.get("current_seed_class", ""),
                "review_bucket": "expanded_auto_curation",
                "suggested_identity_scope": scope,
                "suggested_action": "Auto-classified from European source supplement; keep strict only for exact non-pool named metabolites with stable identifiers.",
                "component_name": reviewed_name,
                "relationship_to_trait": "source_supplement_identity" if scope == "strict_identity" else scope,
                "candidate_names": decision.get("candidate_names", ""),
                "mapped_names": reviewed_name,
                "pubchem_cids": decision.get("reviewed_pubchem_cid", ""),
                "pubchem_titles": reviewed_name,
                "review_status": decision.get("review_status", ""),
                "identity_scope": scope,
                "reviewed_name": reviewed_name,
                "reviewed_pubchem_cid": decision.get("reviewed_pubchem_cid", ""),
                "reviewed_hmdb_id": decision.get("reviewed_hmdb_id", ""),
                "reviewed_chebi_id": decision.get("reviewed_chebi_id", ""),
                "reviewed_inchikey": decision.get("reviewed_inchikey", ""),
                "evidence_source": decision.get("evidence_source", ""),
                "evidence_url": decision.get("evidence_url", ""),
                "notes": f"expanded_pipeline: {decision.get('decision_reason', '')}; super_pathway={decision.get('supplement_super_pathway', '')}; sub_pathway={decision.get('supplement_sub_pathway', '')}",
                "pubmed_id": "36635386",
                "paper_title": "Genomic atlas of the plasma metabolome prioritizes metabolites implicated in human diseases.",
                "autofill_decision": f"expanded_{action}",
                "autofill_confidence": decision.get("decision_confidence", ""),
                "local_cid_status": decision.get("local_cid_status", ""),
            }
        )
    write_csv_rows(review_path, review_rows, fieldnames)
    return {"review_rows": len(review_rows), "added": added, "updated": updated, "skipped": skipped}


def update_compound_overlay(overlay_path: Path, decisions: list[dict[str, Any]], local_cids: set[str], release_id: str) -> dict[str, int]:
    existing_rows, fieldnames = read_csv_rows(overlay_path)
    required_fields = [
        "metabolite_uid",
        "canonical_name",
        "synonyms",
        "pubchem_cid",
        "inchikey",
        "formula",
        "exact_mass",
        "iupac_name",
        "hmdb_id",
        "chebi_id",
        "kegg_id",
        "source_name",
        "source_record_id",
        "source_release",
        "license_id",
        "evidence_source",
        "evidence_url",
        "accession_ids",
        "notes",
    ]
    for field in required_fields:
        if field not in fieldnames:
            fieldnames.append(field)
    by_cid = {normalize_id_token(row.get("pubchem_cid", "")): row for row in existing_rows if row.get("pubchem_cid")}
    added = updated = skipped = 0
    for decision in decisions:
        if decision.get("upgrade_action") != "promote_strict_identity":
            skipped += 1
            continue
        cid = normalize_id_token(decision.get("reviewed_pubchem_cid", ""))
        if not cid or cid in local_cids:
            skipped += 1
            continue
        row = by_cid.get(cid)
        if row is None:
            row = {field: "" for field in fieldnames}
            existing_rows.append(row)
            by_cid[cid] = row
            added += 1
        else:
            updated += 1
        accession_ids = split_values(row.get("accession_ids")) + [decision.get("accession_id", "")]
        row.update(
            {
                "metabolite_uid": row.get("metabolite_uid") or f"manual_pubchem_cid_{cid}",
                "canonical_name": decision.get("reviewed_name", ""),
                "synonyms": join_values([row.get("synonyms", ""), decision.get("candidate_names", ""), decision.get("reported_trait", "")]),
                "pubchem_cid": cid,
                "inchikey": decision.get("reviewed_inchikey", ""),
                "formula": row.get("formula", ""),
                "exact_mass": row.get("exact_mass", ""),
                "iupac_name": row.get("iupac_name", ""),
                "hmdb_id": decision.get("reviewed_hmdb_id", ""),
                "chebi_id": decision.get("reviewed_chebi_id", ""),
                "kegg_id": decision.get("reviewed_kegg_id", ""),
                "source_name": "European_expanded_supplement_review",
                "source_record_id": f"{decision.get('accession_id', '')}:{cid}",
                "source_release": release_id,
                "license_id": "local:european_expanded_identity_review",
                "evidence_source": decision.get("evidence_source", ""),
                "evidence_url": decision.get("evidence_url", ""),
                "accession_ids": join_values(accession_ids),
                "notes": decision.get("decision_reason", ""),
            }
        )
    write_csv_rows(overlay_path, existing_rows, fieldnames)
    return {"overlay_rows": len(existing_rows), "added": added, "updated": updated, "skipped": skipped}


def run_pipeline(args: argparse.Namespace) -> dict[str, Any]:
    workspace = Path(args.workspace).resolve()
    release_id = args.release_id
    supplement_path = (workspace / args.supplement_path).resolve()
    download_supplement(supplement_path, args.supplement_url)
    supplement = load_supplement(supplement_path)
    records = load_european_records(workspace / "raw_lake" / "European" / "European.csv")
    service = build_service(workspace, release_id)
    precheck = service.precheck_metabolites(records)
    candidate_rows, grouped = collect_queue(precheck)
    local_cids = local_pubchem_cids(workspace / "compound_match_index" / release_id)
    decisions = [
        decide_for_group(group, supplement, args.supplement_url, local_cids)
        for group in sorted(grouped.values(), key=lambda row: row.get("accession_id", ""))
    ]

    output_dir = (workspace / args.output_dir).resolve()
    queue_path = output_dir / "european_expanded_review_queue.csv"
    candidate_path = output_dir / "european_expanded_candidate_evidence.csv"
    decisions_path = output_dir / "european_expanded_decisions.csv"
    queue_rows = []
    for group in sorted(grouped.values(), key=lambda row: row.get("accession_id", "")):
        queue_rows.append(
            {
                "original_input_id": group.get("original_input_id", ""),
                "accession_id": group.get("accession_id", ""),
                "reported_trait": group.get("reported_trait", ""),
                "current_seed_class": group.get("current_seed_class", ""),
                "current_reason": group.get("current_reason", ""),
                "source_status": group.get("source_status", ""),
                "top_score": group.get("top_score", ""),
                "top_margin": group.get("top_margin", ""),
                "candidate_count": group.get("candidate_count", ""),
                "candidate_names": group.get("candidate_names", ""),
                "candidate_uids": group.get("candidate_uids", ""),
                "candidate_xrefs": group.get("candidate_xrefs", ""),
            }
        )
    write_csv_rows(queue_path, queue_rows, list(queue_rows[0].keys()) if queue_rows else [])
    write_csv_rows(candidate_path, candidate_rows, list(candidate_rows[0].keys()) if candidate_rows else [])
    write_csv_rows(decisions_path, decisions, list(decisions[0].keys()) if decisions else [])

    review_stats = {"review_rows": 0, "added": 0, "updated": 0, "skipped": len(decisions)}
    overlay_stats = {"overlay_rows": 0, "added": 0, "updated": 0, "skipped": len(decisions)}
    if args.apply:
        review_stats = update_identity_review((workspace / args.review_path).resolve(), decisions)
        overlay_path = (workspace / args.overlay_root / release_id / "compound_identity_overlay.csv").resolve()
        overlay_stats = update_compound_overlay(overlay_path, decisions, local_cids, release_id)

    summary = {
        "release_id": release_id,
        "supplement_path": str(supplement_path),
        "supplement_counts": {"ST2_metabolites": supplement["st2_count"], "ST13_ratios": supplement["st13_count"]},
        "precheck_summary": precheck.get("summary", {}),
        "expanded_summary": precheck.get("expanded_summary", {}),
        "queue_path": str(queue_path),
        "candidate_evidence_path": str(candidate_path),
        "decisions_path": str(decisions_path),
        "decision_counts": dict(Counter(row["decision_class"] for row in decisions)),
        "upgrade_action_counts": dict(Counter(row["upgrade_action"] for row in decisions)),
        "review_update": review_stats,
        "overlay_update": overlay_stats,
    }
    summary_path = output_dir / "european_expanded_pipeline_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return summary


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export, classify, and apply European expanded seed evidence decisions.")
    parser.add_argument("--workspace", default=".")
    parser.add_argument("--release-id", default=DEFAULT_RELEASE_ID)
    parser.add_argument("--supplement-url", default=DEFAULT_SUPPLEMENT_URL)
    parser.add_argument("--supplement-path", default=DEFAULT_SUPPLEMENT_PATH)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--review-path", default=DEFAULT_REVIEW)
    parser.add_argument("--overlay-root", default=DEFAULT_OVERLAY_ROOT)
    parser.add_argument("--apply", action="store_true")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    summary = run_pipeline(args)
    print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
